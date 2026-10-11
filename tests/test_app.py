from datetime import date, datetime, timezone

from publicacoes_app.planilha import (
    ajustar_formula,
    coluna_letra,
    de_serial,
    escolher_aba,
    linha_de_insercao,
    para_serial,
)
from publicacoes_app.prazos import (
    CalendarioForense,
    calcular_fatal,
    calendario_para,
    datas_internas,
    normalizar_cnj,
)

RS = CalendarioForense(uf="RS")


def test_cnj():
    assert normalizar_cnj("CS 5149190-50.2024.8.21.0001 - Ótica") == "5149190-50.2024.8.21.0001"
    assert normalizar_cnj("52112058420268210001") == "5211205-84.2026.8.21.0001"
    assert normalizar_cnj("sem número") is None


def test_disponibilizacao_vira_publicacao_no_proximo_dia_util():
    # WebJur: disponibilizado sexta 09/10/2026; 12/10 é feriado -> publicação 13/10.
    assert RS.proximo_util(date(2026, 10, 9)) == date(2026, 10, 13)


def test_eproc_inicio_da_contagem():
    # Prazo de 5 dias com contagem iniciada em 07/10/2026 (caso Lécio x Ótica Gravataí).
    fatal = calcular_fatal(RS, "inicio_contagem", date(2026, 10, 7), 5)
    assert fatal == date(2026, 10, 14)


def test_publicacao_exclui_dia_do_comeco_e_pula_feriados():
    # Publicado em 13/10; 15 dias úteis; 02/11 é Finados.
    fatal = calcular_fatal(RS, "publicacao", date(2026, 10, 13), 15)
    assert fatal == date(2026, 11, 4)
    # Pela disponibilização, o resultado tem de ser o mesmo.
    assert calcular_fatal(RS, "disponibilizacao", date(2026, 10, 9), 15) == fatal


def test_recesso_forense_suspende_prazo():
    fatal = calcular_fatal(RS, "publicacao", date(2026, 12, 18), 5)
    assert fatal == date(2027, 1, 27)  # 21, 22, 25, 26, 27/01


def test_feriados_justica_federal():
    federal = calendario_para("5013571-67.2025.4.04.7100", "RS", [])
    estadual = calendario_para("5149190-50.2024.8.21.0001", "RS", [])
    quinta_santa = date(2026, 4, 2)
    assert not federal.eh_util(quinta_santa)
    assert estadual.eh_util(quinta_santa)
    assert not estadual.eh_util(date(2026, 6, 4))  # Corpus Christi


def test_prazo_corrido_prorroga_para_dia_util():
    # 10 dias corridos a partir de 01/10 (qui) -> termina 11/10 (dom) -> prorroga p/ 13/10.
    assert calcular_fatal(RS, "publicacao", date(2026, 10, 1), 10, uteis=False) == date(2026, 10, 13)


def test_datas_internas():
    d = datas_internas(RS, date(2026, 10, 20), antecedencia_lo=3, antecedencia_revisao=1)
    assert d.lo == date(2026, 10, 15)
    assert d.revisao == date(2026, 10, 19)
    # Prazo curto: sugestões não ficam no passado.
    d = datas_internas(RS, date(2026, 10, 14), 3, 1, hoje=date(2026, 10, 11))
    assert d.lo == date(2026, 10, 13) and d.revisao == date(2026, 10, 13)


def test_escolher_aba():
    abas = ["W41 13102026", "W40 05102026", "W39 28092026", "Contencioso"]
    assert escolher_aba(abas, date(2026, 10, 7)) == "W40 05102026"
    # Domingo: já considera a semana seguinte.
    assert escolher_aba(abas, date(2026, 10, 11)) == "W41 13102026"
    assert escolher_aba(abas, date(2026, 10, 16)) == "W41 13102026"
    # Semana sem aba ainda: usa a mais recente.
    assert escolher_aba(abas, date(2026, 10, 21)) == "W41 13102026"


def test_serial_excel():
    assert para_serial(date(2026, 10, 13)) == 46308
    assert de_serial(46308) == date(2026, 10, 13)
    assert de_serial("13/10/2026") == date(2026, 10, 13)
    assert de_serial("sine die") is None
    assert coluna_letra(12) == "M" and coluna_letra(26) == "AA"


def test_linha_de_insercao_ordena_por_prazo_lo():
    s = para_serial
    linhas = [
        ["Cliente", "Prazo LO"],
        ["A", s(date(2026, 10, 9))],
        ["B", s(date(2026, 10, 13))],
        ["C", s(date(2026, 10, 15))],
        ["D", "sine die"],
        ["E", ""],
    ]
    assert linha_de_insercao(linhas, 1, date(2026, 10, 13)) == 3
    assert linha_de_insercao(linhas, 1, date(2026, 10, 1)) == 1
    assert linha_de_insercao(linhas, 1, date(2026, 12, 1)) == 4
    assert linha_de_insercao(linhas, 1, None) == 6


def test_ajustar_formula():
    f = '=IF(K2="Concluído","",IF(AND(ISNUMBER(G2),NETWORKDAYS(TODAY(),G2)-1=2),"D-2",IF(E2<TODAY(),"VENCIDO","")))'
    novo = ajustar_formula(f, 2, 17)
    assert "K17" in novo and "G17" in novo and "E17" in novo
    assert "2)" not in novo.replace("=2", "")
    assert '"D-2"' in novo
    assert ajustar_formula("=LOG10(A10)", 10, 11) == "=LOG10(A11)"


# ------------------------------------------------------------ Legal One

from datetime import datetime

from openpyxl import Workbook, load_workbook

from publicacoes_app import atribuicoes
from publicacoes_app.legalone import agrupar, ler_tarefas, localizar_arquivo


def _planilha_lo(caminho, cabecalho, linhas):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tarefas Pendentes"
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
    wb.save(caminho)
    return str(caminho)


NOVO = ["ID", "Status", "Vencimento", "Responsável", "Cliente", "Processo", "Pasta", "Tipo", "Tarefa"]


def test_le_formato_novo_e_agrupa_executor_revisor(tmp_path):
    arq = _planilha_lo(tmp_path / "Tarefas_Pendentes_LegalOne_Outubro-2026.xlsx", NOVO, [
        [101, "A Vencer", datetime(2026, 10, 13), "Tiago Coster", "Lécio José de Barros Filho",
         "5149190-50.2024.8.21.0001", "1234", "Prazo", "Contrarrazões aos Embargos de Declaração"],
        [102, "A Vencer", datetime(2026, 10, 15), "Tiago Coster", "Lécio José de Barros Filho",
         "5149190-50.2024.8.21.0001", "1234", "Revisão", "Revisar contrarrazões aos EDs"],
        [103, "A Vencer", datetime(2026, 10, 20), "Controladoria", "AIG",
         "5004654-79.2025.4.03.6100", "88", "Prazo", "RESP/RE"],
    ])
    tarefas = ler_tarefas(arq)
    assert [t.id for t in tarefas] == ["101", "102", "103"]
    assert tarefas[1].de_revisao and not tarefas[0].de_revisao

    lecio, aig = agrupar(tarefas, RS)
    assert lecio.executor.id == "101" and lecio.revisor.id == "102"
    assert (lecio.lo, lecio.revisao, lecio.fatal) == (date(2026, 10, 13), date(2026, 10, 15), date(2026, 10, 19))
    assert aig.revisor is None and aig.fatal == date(2026, 10, 22)
    assert aig.avisos


def test_le_formato_antigo_cliente_processo_juntos(tmp_path):
    arq = _planilha_lo(tmp_path / "x.xlsx", ["Status", "Vencimento", "Responsável", "Cliente/Processo", "Tarefa"], [
        ["Vencido", "07/10/2026", "Pedro Delapieve", "Hyva do Brasil - 5056840-72.2026.8.21.0001", "Verificar custas"],
    ])
    [t] = ler_tarefas(arq)
    assert t.cliente == "Hyva do Brasil"
    assert t.processo == "5056840-72.2026.8.21.0001"
    assert t.vencimento == date(2026, 10, 7)
    assert t.id.startswith("h")  # sem coluna ID: identificador estável derivado do conteúdo
    assert ler_tarefas(arq)[0].id == t.id


def test_mesmo_processo_dois_prazos_nao_se_misturam():
    from publicacoes_app.legalone import TarefaLO

    p = "5000561-81.2023.4.04.0000"
    tarefas = [
        TarefaLO("1", date(2026, 10, 14), processo=p, descricao="RESP"),
        TarefaLO("2", date(2026, 10, 16), processo=p, descricao="Revisão RESP"),
        TarefaLO("3", date(2026, 11, 20), processo=p, descricao="Memoriais"),
    ]
    d1, d2 = agrupar(tarefas, RS)
    assert d1.ids == ["1", "2"] and d2.ids == ["3"]


def test_localizar_arquivo_mais_recente(tmp_path):
    import os, time

    a = _planilha_lo(tmp_path / "Tarefas_Pendentes_LegalOne_Setembro-2026.xlsx", NOVO, [])
    b = _planilha_lo(tmp_path / "Tarefas_Pendentes_LegalOne_Outubro-2026.xlsx", NOVO, [])
    os.utime(a, (time.time() - 100, time.time() - 100))
    assert localizar_arquivo(str(tmp_path), "Tarefas_Pendentes_LegalOne_*.xlsx") == b


def test_planilha_de_atribuicoes(tmp_path):
    caminho = atribuicoes.caminho_do_dia(str(tmp_path), date(2026, 10, 13))
    a = atribuicoes.Atribuicao("Lécio", "5149190-50.2024.8.21.0001", "1234", "Contrarrazões",
                               "Danieli", date(2026, 10, 13), "Pedro", date(2026, 10, 15), date(2026, 10, 19),
                               id_executor="101", responsavel_executor="Tiago Coster")
    assert atribuicoes.registrar(caminho, a) == 2
    assert atribuicoes.registrar(caminho, a) == 3
    ws = load_workbook(caminho)["Atribuições"]
    cab = [c.value for c in ws[1]]
    linha = dict(zip(cab, [c.value for c in ws[2]]))
    assert linha["Executor"] == "Danieli" and linha["ID tarefa executor (L.O.)"] == "101"
    assert linha["Andamento lançado"] == "Não"
    assert linha["Texto do andamento"] == "PEDRO - Novo Fluxo: Executor: Danieli, Revisor: Pedro."
    a.revisor = ""
    assert a.texto() == "PEDRO - Novo Fluxo: Executor: Danieli."


# ------------------------------------------------------------ fluxo completo

class _FakeGraph:
    def __init__(self):
        self.posts = []

    def post(self, caminho, json=None, **kw):
        self.posts.append((caminho, json))
        return {"webLink": "https://outlook/evento"}


class _FakePlanilha:
    def __init__(self, duplicada=False):
        self.incluidas = []
        self._dup = duplicada

    def buscar_processo(self, abas, processo):
        return {"Cliente": "Lécio José de Barros Filho", "Natureza": "Contencioso",
                "Caso": "CS 5149190-50.2024.8.21.0001 - Ótica Gravataí"}

    def duplicada(self, aba, processo, fatal):
        return self._dup

    def incluir(self, aba, tarefa):
        self.incluidas.append((aba, tarefa))
        return 10


def _demanda_lecio():
    from publicacoes_app.legalone import TarefaLO

    p = "5149190-50.2024.8.21.0001"
    return agrupar([
        TarefaLO("101", date(2026, 10, 13), "Tiago Coster", "Lécio", p, "1234", "Prazo", "Contrarrazões aos EDs"),
        TarefaLO("102", date(2026, 10, 15), "Tiago Coster", "Lécio", p, "1234", "Revisão", "Revisar contrarrazões"),
    ], RS)[0]


def test_fluxo_interativo(monkeypatch, tmp_path):
    from publicacoes_app import cli
    from publicacoes_app.config import Config

    cfg = Config(client_id="x", tenant_id="y", uf="RS", equipe={"Pedro": "p@x", "Danieli": ""},
                 eventos=["lo", "fatal"], base_dir=tmp_path)
    respostas = iter([
        "A",            # atribuir
        "2",            # executor: Danieli
        "1",            # revisor: Pedro
        "", "", "",     # fatal, LO e revisão vindos do L.O.
        "", "", "",     # cliente / natureza / caso vindos da Programação
        "Contrarrazões aos Embargos de Declaração",
        "", "",         # prioridade / observações
        "s",            # confirmar
    ])
    monkeypatch.setattr("builtins.input", lambda _="": next(respostas))
    monkeypatch.setattr(cli, "hoje", lambda: date(2026, 10, 9))
    g, plan = _FakeGraph(), _FakePlanilha()
    estado = cli.Estado(tmp_path / "estado.json")
    caminho = str(tmp_path / "Atribuicoes.xlsx")
    d = _demanda_lecio()

    r = cli.processar_demanda(1, 1, d, cfg, g, plan, "W41 13102026", [], estado, caminho, simular=False)

    assert r == "atribuida"
    [(aba, t)] = plan.incluidas
    assert (t.executor, t.revisor) == ("Danieli", "Pedro")
    assert (t.lo, t.revisao, t.fatal) == (date(2026, 10, 13), date(2026, 10, 15), date(2026, 10, 19))
    assert t.cliente == "Lécio José de Barros Filho"
    eventos = [j for c, j in g.posts if c == "/me/events"]
    assert [e["start"]["dateTime"][:10] for e in eventos] == ["2026-10-13", "2026-10-19"]
    assert estado.tratada(d)
    ws = load_workbook(caminho)["Atribuições"]
    assert ws.cell(row=2, column=6).value == "Danieli"


def test_fluxo_ja_consta_sugere_ignorar(monkeypatch, tmp_path):
    from publicacoes_app import cli
    from publicacoes_app.config import Config

    cfg = Config(client_id="x", tenant_id="y", equipe={"Pedro": ""}, base_dir=tmp_path)
    monkeypatch.setattr("builtins.input", lambda _="": "")  # Enter = padrão (Ignorar)
    estado = cli.Estado(tmp_path / "estado.json")
    d = _demanda_lecio()
    r = cli.processar_demanda(1, 1, d, cfg, _FakeGraph(), _FakePlanilha(duplicada=True), "W41", [],
                              estado, str(tmp_path / "a.xlsx"), simular=False)
    assert r == "ignorada" and estado.tratada(d)
