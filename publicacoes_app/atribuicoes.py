"""Planilha diária com as atribuições feitas (base para os andamentos no L.O.)."""

from __future__ import annotations

import os
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ABA = "Atribuições"
COLUNAS = [
    ("Registrado em", 17),
    ("Cliente", 22),
    ("Processo", 27),
    ("Pasta L.O.", 12),
    ("Providência", 34),
    ("Executor", 12),
    ("Prazo LO", 11),
    ("Revisor", 12),
    ("Revisão", 11),
    ("Fatal", 11),
    ("ID tarefa executor (L.O.)", 14),
    ("Responsável atual executor (L.O.)", 18),
    ("ID tarefa revisor (L.O.)", 14),
    ("Responsável atual revisor (L.O.)", 18),
    ("Texto do andamento", 70),
    ("Andamento lançado", 11),
]
COLUNAS_DATA = {"Prazo LO", "Revisão", "Fatal"}

# Padrão usado no L.O. para ajuste de fluxo (andamento do tipo "Andamento").
MODELO_PADRAO = "{autor} - Novo Fluxo: Executor: {executor}, Revisor: {revisor}."
MODELO_SEM_REVISOR = "{autor} - Novo Fluxo: Executor: {executor}."


@dataclass
class Atribuicao:
    cliente: str
    processo: str
    pasta: str
    providencia: str
    executor: str
    lo: date | None
    revisor: str
    revisao: date | None
    fatal: date | None
    id_executor: str = ""
    responsavel_executor: str = ""
    id_revisor: str = ""
    responsavel_revisor: str = ""

    def texto(self, modelo: str = MODELO_PADRAO, autor: str = "PEDRO") -> str:
        fmt = lambda d: f"{d:%d/%m/%Y}" if d else "não definido"
        if not self.revisor and modelo == MODELO_PADRAO:
            modelo = MODELO_SEM_REVISOR
        return modelo.format(
            autor=autor,
            providencia=self.providencia or "-",
            executor=self.executor or "não definido",
            revisor=self.revisor or "sem revisor",
            lo=fmt(self.lo),
            revisao=fmt(self.revisao),
            fatal=fmt(self.fatal),
            cliente=self.cliente,
            processo=self.processo,
        )


def caminho_do_dia(pasta: str, dia: date) -> str:
    return os.path.join(os.path.expanduser(pasta), f"Atribuicoes_{dia:%Y-%m-%d}.xlsx")


def _novo_livro() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = ABA
    ws.append([c for c, _ in COLUNAS])
    for i, (_, largura) in enumerate(COLUNAS, 1):
        ws.column_dimensions[get_column_letter(i)].width = largura
        cel = ws.cell(row=1, column=i)
        cel.font = Font(name="Arial", bold=True, color="FFFFFF")
        cel.fill = PatternFill("solid", fgColor="00214D")
        cel.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}1"
    return wb


def registrar(caminho: str, a: Atribuicao, modelo: str = MODELO_PADRAO, autor: str = "PEDRO") -> int:
    """Acrescenta a atribuição à planilha do dia (cria se não existir). Retorna a linha."""
    wb = load_workbook(caminho) if os.path.exists(caminho) else _novo_livro()
    ws = wb[ABA]
    valores = {
        "Registrado em": datetime.now().replace(microsecond=0),
        "Cliente": a.cliente,
        "Processo": a.processo,
        "Pasta L.O.": a.pasta,
        "Providência": a.providencia,
        "Executor": a.executor,
        "Prazo LO": a.lo,
        "Revisor": a.revisor,
        "Revisão": a.revisao,
        "Fatal": a.fatal,
        "ID tarefa executor (L.O.)": a.id_executor,
        "Responsável atual executor (L.O.)": a.responsavel_executor,
        "ID tarefa revisor (L.O.)": a.id_revisor,
        "Responsável atual revisor (L.O.)": a.responsavel_revisor,
        "Texto do andamento": a.texto(modelo, autor),
        "Andamento lançado": "Não",
    }
    linha = ws.max_row + 1
    for i, (nome, _) in enumerate(COLUNAS, 1):
        cel = ws.cell(row=linha, column=i, value=valores[nome])
        cel.font = Font(name="Arial")
        cel.alignment = Alignment(wrap_text=nome in ("Texto do andamento", "Providência"), vertical="top")
        if nome in COLUNAS_DATA:
            cel.number_format = "DD/MM/YYYY"
        elif nome == "Registrado em":
            cel.number_format = "DD/MM/YYYY HH:MM"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUNAS))}{linha}"
    os.makedirs(os.path.dirname(caminho) or ".", exist_ok=True)
    wb.save(caminho)
    return linha
