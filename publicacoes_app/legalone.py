"""Leitura da planilha gerada pela rotina "Prazos do Mês no L.O."

Arquivo: ``Tarefas_Pendentes_LegalOne_<Mês>-<Ano>.xlsx``, aba "Tarefas Pendentes".
Formato novo (colunas separadas): ID | Status | Vencimento | Responsável |
Cliente | Processo | Pasta | Tipo | Tarefa. O formato antigo (Status |
Vencimento | Responsável | Cliente/Processo | Tarefa) também é aceito.

No L.O. a controladoria agenda uma tarefa para o executor e outra para o
revisor; a data de conclusão da tarefa do revisor é o D-2 do fatal.
"""

from __future__ import annotations

import glob
import hashlib
import os
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from .prazos import CalendarioForense, normalizar_cnj

ABA = "Tarefas Pendentes"
REVISAO_RE = re.compile(r"revis", re.I)
# Distância máxima entre a tarefa do executor e a do revisor de um mesmo prazo.
JANELA_PAREAMENTO = timedelta(days=15)


def _chave(texto) -> str:
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", t.casefold()).strip()


ALIASES = {
    "id": ("id", "id l o", "id lo", "codigo"),
    "status": ("status",),
    "vencimento": ("vencimento", "data conclusao prevista", "conclusao prevista", "data"),
    "responsavel": ("responsavel",),
    "cliente": ("cliente", "cliente principal"),
    "processo": ("processo", "numero do processo"),
    "pasta": ("pasta",),
    "tipo": ("tipo", "subtipo", "tipo subtipo"),
    "tarefa": ("tarefa", "descricao"),
    "cliente_processo": ("cliente processo", "vinculo"),
}


@dataclass
class TarefaLO:
    id: str
    vencimento: date
    responsavel: str = ""
    cliente: str = ""
    processo: str | None = None
    pasta: str = ""
    tipo: str = ""
    descricao: str = ""
    status: str = ""

    @property
    def de_revisao(self) -> bool:
        return bool(REVISAO_RE.search(f"{self.tipo} {self.descricao}"))


@dataclass
class Demanda:
    """Um prazo: tarefa do executor e, se houver, a do revisor."""

    executor: TarefaLO | None
    revisor: TarefaLO | None
    fatal: date
    avisos: list[str] = field(default_factory=list)

    @property
    def tarefas(self) -> list[TarefaLO]:
        return [t for t in (self.executor, self.revisor) if t]

    @property
    def principal(self) -> TarefaLO:
        return self.executor or self.revisor

    @property
    def ids(self) -> list[str]:
        return [t.id for t in self.tarefas]

    @property
    def lo(self) -> date | None:
        return self.executor.vencimento if self.executor else None

    @property
    def revisao(self) -> date | None:
        return self.revisor.vencimento if self.revisor else None


def localizar_arquivo(pasta: str, padrao: str) -> str | None:
    """Arquivo mais recente que casa com o padrão na pasta."""
    arquivos = glob.glob(os.path.join(os.path.expanduser(pasta), padrao))
    return max(arquivos, key=os.path.getmtime) if arquivos else None


def _como_data(valor) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y"):
            try:
                return datetime.strptime(valor.strip()[:10], fmt).date()
            except ValueError:
                continue
    return None


def _separar_cliente_processo(texto: str) -> tuple[str, str | None, str]:
    """'Cliente X - 5000000-00.2024.4.04.7100 (Pasta 123)' -> (cliente, processo, pasta)."""
    processo = normalizar_cnj(texto)
    pasta = ""
    if m := re.search(r"(?i)pasta[:\s#nº]*([\w./-]+)", texto):
        pasta = m.group(1)
    cliente = texto
    if processo:
        cliente = re.split(r"[\d]{7}-?\d{2}\.?\d{4}", texto)[0]
    cliente = re.sub(r"(?i)pasta[:\s#nº]*[\w./-]+", "", cliente)
    cliente = re.sub(r"[\s\-–|/,;:()]+$", "", cliente.strip())
    return cliente.strip(), processo, pasta


def _id_sintetico(*partes) -> str:
    return "h" + hashlib.sha1("|".join(map(str, partes)).encode()).hexdigest()[:12]


def ler_tarefas(caminho: str) -> list[TarefaLO]:
    wb = load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[ABA] if ABA in wb.sheetnames else wb.worksheets[0]
    linhas = ws.iter_rows(values_only=True)
    cabecalho = [_chave(c) for c in next(linhas, [])]

    def indice(campo: str) -> int | None:
        for alias in ALIASES[campo]:
            if alias in cabecalho:
                return cabecalho.index(alias)
        return None

    idx = {campo: indice(campo) for campo in ALIASES}
    if idx["vencimento"] is None:
        raise ValueError(f"Coluna de vencimento não encontrada em {caminho}: {cabecalho}")

    def pega(linha, campo) -> str:
        i = idx[campo]
        v = linha[i] if i is not None and i < len(linha) else None
        return "" if v is None else str(v).strip()

    tarefas = []
    for linha in linhas:
        venc = _como_data(linha[idx["vencimento"]] if idx["vencimento"] < len(linha) else None)
        if not venc:
            continue
        cliente, processo, pasta = pega(linha, "cliente"), pega(linha, "processo"), pega(linha, "pasta")
        if idx["cliente_processo"] is not None:
            c2, p2, pa2 = _separar_cliente_processo(pega(linha, "cliente_processo"))
            cliente, processo, pasta = cliente or c2, processo or p2, pasta or pa2
        processo = normalizar_cnj(processo or "") or (processo or None)
        descricao = pega(linha, "tarefa")
        responsavel = pega(linha, "responsavel")
        tarefas.append(
            TarefaLO(
                id=pega(linha, "id") or _id_sintetico(cliente, processo, descricao, venc),
                vencimento=venc,
                responsavel=responsavel,
                cliente=cliente,
                processo=processo,
                pasta=pasta,
                tipo=pega(linha, "tipo"),
                descricao=descricao,
                status=pega(linha, "status"),
            )
        )
    wb.close()
    return tarefas


def _vinculo(t: TarefaLO) -> str:
    return re.sub(r"\D", "", t.processo or "") or _chave(t.pasta) or _chave(t.cliente)


def agrupar(tarefas: list[TarefaLO], cal: CalendarioForense, d_menos: int = 2) -> list[Demanda]:
    """Junta, por processo/pasta, a tarefa do executor com a do revisor do mesmo prazo.

    Fatal sugerido = vencimento da tarefa do revisor (ou, sem ela, a do
    executor) + ``d_menos`` dias úteis.
    """
    por_vinculo: dict[str, list[TarefaLO]] = {}
    for t in sorted(tarefas, key=lambda t: t.vencimento):
        por_vinculo.setdefault(_vinculo(t), []).append(t)

    demandas = []
    for grupo in por_vinculo.values():
        revisoes = [t for t in grupo if t.de_revisao]
        for exe in (t for t in grupo if not t.de_revisao):
            par = next(
                (r for r in revisoes if exe.vencimento <= r.vencimento <= exe.vencimento + JANELA_PAREAMENTO),
                None,
            )
            if par:
                revisoes.remove(par)
            demandas.append(_demanda(exe, par, cal, d_menos))
        demandas.extend(_demanda(None, r, cal, d_menos) for r in revisoes)
    return sorted(demandas, key=lambda d: (d.fatal, d.principal.cliente))


def _demanda(exe: TarefaLO | None, rev: TarefaLO | None, cal: CalendarioForense, d_menos: int) -> Demanda:
    base = (rev or exe).vencimento
    d = Demanda(executor=exe, revisor=rev, fatal=cal.somar_uteis(base, d_menos))
    if not rev:
        d.avisos.append("Sem tarefa de revisão no L.O.; fatal estimado a partir da tarefa do executor.")
    if not exe:
        d.avisos.append("Só há tarefa de revisão no L.O. para este prazo.")
    return d
